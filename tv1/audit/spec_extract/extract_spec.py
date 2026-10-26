import openpyxl

path = r"C:\Users\phucv\Downloads\csd\Phan_cong_chuan_theo_de_Sorting_pdqsort.xlsx"
wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
print("SHEETS=" + "|".join(wb.sheetnames))
for wanted in ("README", "Tổng quan", "Ma trận RQ", "TV1", "Quy ước chung"):
    if wanted not in wb.sheetnames:
        print("MISSING_SHEET=" + wanted)
        continue
    ws = wb[wanted]
    print("BEGIN_SHEET=" + wanted + " max_row=" + str(ws.max_row) + " max_col=" + str(ws.max_column))
    for row_index, row in enumerate(ws.iter_rows(values_only=True), 1):
        if any(value is not None and str(value).strip() != "" for value in row):
            print("ROW " + str(row_index) + " | " + " | ".join(
                "" if value is None else str(value).replace("\n", "\\n") for value in row))
    print("END_SHEET=" + wanted)
