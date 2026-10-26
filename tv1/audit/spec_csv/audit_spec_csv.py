#!/usr/bin/env python3
"""Read-only audit for the TV1 QuickSort CSV and workbook specification."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl


ROOT = Path(r"C:/Users/phucv/Downloads/csd")
AUDIT_DIR = ROOT / "tv1" / "audit" / "spec_csv"
WORKBOOK = ROOT / "Phan_cong_chuan_theo_de_Sorting_pdqsort.xlsx"
CSV_PATH = ROOT / "tv1" / "out" / "quick_sort_results.csv"
SOURCE = ROOT / "tv1" / "src" / "main" / "java" / "sorting" / "bench" / "QuickBench.java"
META_PATHS = [
    ROOT / "tv1" / "out" / "quick_sort_metadata.json",
    ROOT / "tv1" / "out" / "quick_sort_metadata_count.json",
    ROOT / "tv1" / "out" / "quick_sort_metadata_d2.json",
    ROOT / "tv1" / "out" / "quick_sort_metadata_d2_count.json",
    ROOT / "tv1" / "out" / "quick_sort_metadata_50m.json",
]

REQUIRED_MINIMUM = [
    "algorithm",
    "variant",
    "dataset",
    "N",
    "K",
    "seed",
    "elapsed_ms",
    "heap_before_bytes",
    "heap_peak_bytes",
    "output_check",
]
COUNT_FIELDS = ["comparisons", "max_depth"]
EXPECTED_CODE_HEADER = [
    "algorithm",
    "variant",
    "dataset",
    "N",
    "K",
    "seed",
    "elapsed_ms",
    "heap_before_bytes",
    "heap_peak_bytes",
    "output_check",
    "comparisons",
    "max_depth",
    "mode",
    "aux_bytes",
    "sorted_ok",
    "rep",
    "jdk",
    "os",
    "arch",
    "cores",
    "max_heap_bytes",
    "run_id",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest().upper()


def workbook_cells(sheet_name: str) -> dict[str, object]:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    if sheet_name not in wb.sheetnames:
        raise RuntimeError(f"missing workbook sheet: {sheet_name}")
    ws = wb[sheet_name]
    return {cell.coordinate: cell.value for row in ws.iter_rows() for cell in row
            if cell.value is not None}


def parse_csv() -> tuple[list[str], list[dict[str, str]]]:
    with CSV_PATH.open("r", encoding="utf-8", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None:
            raise RuntimeError("CSV has no header")
        rows = list(reader)
        return list(reader.fieldnames), rows


def val(row: dict[str, str], field: str) -> str:
    return (row.get(field) or "").strip()


def parse_code_header() -> list[str]:
    text = SOURCE.read_text(encoding="utf-8")
    match = re.search(r'CSV_HEADER\s*=\s*(.*?);', text, re.S)
    if not match:
        raise RuntimeError("CSV_HEADER not found in QuickBench.java")
    parts = re.findall(r'"([^"]*)"', match.group(1))
    return "".join(parts).split(",")


def parse_matrix_ns() -> dict[str, object]:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    ws = wb["Ma trận RQ"]
    text = str(ws["E3"].value or "")
    numbers = []
    for token in re.split(r";", text):
        token = token.strip()
        if not token:
            continue
        if token.startswith("thêm"):
            token = token[4:].strip()
        token = token.split("nếu đủ RAM")[0].strip()
        cleaned = token.replace("×", "*").replace("^", "**")
        try:
            value = int(eval(cleaned, {"__builtins__": {}}, {}))  # controlled workbook text
        except Exception:
            continue
        if value not in numbers:
            numbers.append(value)
    return {"text": text, "numbers": numbers,
            "optional_50m": "5×10^7" in text and "nếu đủ RAM" in text}


def metadata_summary() -> list[dict[str, object]]:
    out = []
    for path in META_PATHS:
        data = json.loads(path.read_text(encoding="utf-8"))
        out.append({"path": path.name, "data": data, "sha256": sha256(path)})
    return out


def metadata_alignment(rows: list[dict[str, str]], metadata: list[dict[str, object]]) -> dict[str, object]:
    checks = []
    row_coverage = []
    all_ns_covered = True
    all_fields_match = True
    for entry in metadata:
        data = entry["data"]
        run_id = str(data["run_id"])
        mode = str(data["mode"])
        dataset = str(data["dataset"])
        ns = [int(x) for x in data["ns"]]
        matches = [r for r in rows
                   if val(r, "run_id") == run_id
                   and val(r, "mode") == mode
                   and val(r, "dataset") == dataset]
        covered = sorted({int(val(r, "N")) for r in matches})
        missing = sorted(set(ns) - set(covered))
        extra = sorted(set(covered) - set(ns))
        field_mismatches = []
        for row in matches:
            expected = {
                "seed": str(data["seed"]),
                "jdk": str(data["java_version"]),
                "cores": str(data["available_processors"]),
                "max_heap_bytes": str(data["max_heap_bytes"]),
                "rep": str(data["reps"]),
            }
            for field, wanted in expected.items():
                if val(row, field) != wanted:
                    field_mismatches.append({"line": rows.index(row) + 2, "field": field,
                                              "expected": wanted, "actual": val(row, field)})
        all_ns_covered = all_ns_covered and not missing
        all_fields_match = all_fields_match and not field_mismatches
        checks.append({
            "path": entry["path"], "run_id": run_id, "mode": mode, "dataset": dataset,
            "metadata_ns": ns, "csv_ns": covered, "missing_ns": missing, "extra_ns": extra,
            "field_mismatches": field_mismatches,
        })

    for line, row in enumerate(rows, start=2):
        n = int(val(row, "N"))
        covered_by = [entry["path"] for entry in metadata
                      if str(entry["data"]["run_id"]) == val(row, "run_id")
                      and str(entry["data"]["mode"]) == val(row, "mode")
                      and str(entry["data"]["dataset"]) == val(row, "dataset")
                      and n in [int(x) for x in entry["data"]["ns"]]]
        row_coverage.append({"line": line, "dataset": val(row, "dataset"), "N": n,
                             "mode": val(row, "mode"), "covered_by": covered_by,
                             "covered": bool(covered_by)})
    return {
        "files": checks,
        "row_coverage": row_coverage,
        "all_metadata_ns_covered": all_ns_covered,
        "all_metadata_fields_match": all_fields_match,
        "all_rows_covered_by_at_least_one_metadata": all(x["covered"] for x in row_coverage),
        "all_metadata_share_same_run_id": len({entry["data"]["run_id"] for entry in metadata}) == 1,
    }


def group_rows(rows: list[dict[str, str]], *fields: str) -> dict[tuple[str, ...], list[dict[str, str]]]:
    grouped = defaultdict(list)
    for row in rows:
        grouped[tuple(val(row, f) for f in fields)].append(row)
    return dict(grouped)


def independent_values() -> dict[int, dict[str, int]]:
    """Compile and run our independent Java generator, then parse its output."""
    classes = AUDIT_DIR / "classes"
    classes.mkdir(parents=True, exist_ok=True)
    subprocess.run(
        ["javac", "--release", "8", "-d", str(classes), str(AUDIT_DIR / "IndependentD1.java")],
        check=True, capture_output=True, text=True,
    )
    proc = subprocess.run(
        ["java", "-Xmx4g", "-cp", str(classes), "IndependentD1"],
        check=True, capture_output=True, text=True, timeout=300,
    )
    result: dict[int, dict[str, int]] = {}
    for line in proc.stdout.splitlines():
        m = re.match(r"D1 n=(\d+) raw_output_check=(-?\d+) sorted_output_check=(-?\d+)", line)
        if m:
            n = int(m.group(1))
            result[n] = {
                "raw_output_check": int(m.group(2)),
                "sorted_output_check": int(m.group(3)),
            }
    if not result:
        raise RuntimeError("no independent D1 lines parsed")
    return result


def main() -> int:
    tv1 = workbook_cells("TV1")
    common = workbook_cells("Quy ước chung")
    matrix = parse_matrix_ns()
    header, rows = parse_csv()
    code_header = parse_code_header()
    independent = independent_values()

    result: dict[str, object] = {
        "inputs": {
            "workbook": {"path": str(WORKBOOK), "sha256": sha256(WORKBOOK)},
            "source": {"path": str(SOURCE), "sha256": sha256(SOURCE)},
            "csv": {"path": str(CSV_PATH), "sha256": sha256(CSV_PATH)},
        },
        "workbook": {
            "TV1": tv1,
            "Quy ước chung": common,
            "matrix_rq1_n": matrix,
        },
        "csv": {
            "header": header,
            "row_count_excluding_header": len(rows),
            "field_count_header": len(header),
            "field_count_rows": sorted(set(len(r) for r in rows)),
            "rows": rows,
        },
        "checks": {},
    }

    checks = result["checks"]
    checks["header_exact_code_order"] = header == code_header
    checks["minimum_required_fields_in_order"] = all(
        field in header and header.index(field) == i
        for i, field in enumerate(REQUIRED_MINIMUM)
    )
    checks["count_fields_present"] = all(field in header for field in COUNT_FIELDS)
    checks["csv_header_matches_expected_code_header"] = header == EXPECTED_CODE_HEADER

    d1_rows = [r for r in rows if val(r, "dataset") == "d1"]
    d2_rows = [r for r in rows if val(r, "dataset") == "d2"]
    rq1_ns = matrix["numbers"]
    required_ns = [n for n in rq1_ns if not (matrix["optional_50m"] and n == 50_000_000)]
    checks["matrix"] = {
        "rq1_text": matrix["text"],
        "parsed_ns": rq1_ns,
        "mandatory_ns": required_ns,
        "optional_50m": matrix["optional_50m"],
        "expected_quick_rq1_rows_minimum": 2 * len(required_ns),
        "expected_quick_rq1_rows_with_optional_50m": 2 * len(rq1_ns),
        "delivered_d1_row_count": len(d1_rows),
        "delivered_d1_time_rows": len([r for r in d1_rows if val(r, "mode") == "time"]),
        "delivered_d1_count_rows": len([r for r in d1_rows if val(r, "mode") == "count"]),
        "delivered_d1_ns": sorted({int(val(r, "N")) for r in d1_rows}),
        "delivered_d2_row_count": len(d2_rows),
        "delivered_d2_ns": sorted({int(val(r, "N")) for r in d2_rows}),
        "d1_missing_required_time": sorted(set(required_ns) - {int(val(r, "N")) for r in d1_rows if val(r, "mode") == "time"}),
        "d1_missing_required_count": sorted(set(required_ns) - {int(val(r, "N")) for r in d1_rows if val(r, "mode") == "count"}),
    }

    mode_errors = []
    for i, row in enumerate(rows, start=2):
        mode = val(row, "mode")
        if mode == "time":
            bad = []
            if not val(row, "elapsed_ms"):
                bad.append("elapsed_ms_empty")
            if not val(row, "heap_before_bytes") or not val(row, "heap_peak_bytes"):
                bad.append("heap_empty")
            if val(row, "comparisons") or val(row, "max_depth"):
                bad.append("count_fields_nonempty")
            if bad:
                mode_errors.append({"line": i, "mode": mode, "errors": bad, "row": row})
        elif mode == "count":
            bad = []
            if val(row, "elapsed_ms") or val(row, "heap_before_bytes") or val(row, "heap_peak_bytes"):
                bad.append("time_or_heap_nonempty")
            if not val(row, "comparisons") or not val(row, "max_depth"):
                bad.append("count_field_empty")
            if bad:
                mode_errors.append({"line": i, "mode": mode, "errors": bad, "row": row})
        else:
            mode_errors.append({"line": i, "mode": mode, "errors": ["unknown_mode"], "row": row})
    checks["mode_separation"] = {"ok": not mode_errors, "errors": mode_errors}

    sorted_errors = [{"line": i, "sorted_ok": val(r, "sorted_ok")}
                     for i, r in enumerate(rows, start=2) if val(r, "sorted_ok") != "true"]
    checks["sorted_ok_all_true"] = {"ok": not sorted_errors, "errors": sorted_errors}

    output_groups = group_rows(rows, "dataset", "N", "seed")
    output_checks = []
    for key, group in sorted(output_groups.items(), key=lambda kv: (kv[0][0], int(kv[0][1]))):
        values = [val(r, "output_check") for r in group]
        modes = [val(r, "mode") for r in group]
        output_checks.append({
            "dataset": key[0], "N": int(key[1]), "seed": key[2],
            "modes": modes, "output_checks": values,
            "all_equal": len(set(values)) == 1,
            "has_time": "time" in modes, "has_count": "count" in modes,
        })
    checks["output_check_time_count_equality"] = {
        "groups": output_checks,
        "all_groups_equal": all(g["all_equal"] for g in output_checks),
        "groups_with_both_modes": [g for g in output_checks if g["has_time"] and g["has_count"]],
    }

    independent_compare = []
    for row in rows:
        n = int(val(row, "N"))
        expected = independent.get(n, {}).get("sorted_output_check")
        actual = int(val(row, "output_check"))
        independent_compare.append({
            "dataset": val(row, "dataset"), "N": n,
            "mode": val(row, "mode"), "actual": actual,
            "independent_sorted_output_check": expected,
            "match": expected == actual,
        })
    checks["independent_d1_reproducibility"] = {
        "independent_values": independent,
        "comparisons": independent_compare,
        "all_match": all(x["match"] for x in independent_compare),
    }

    duplicate_keys = []
    key_counter = Counter(
        tuple(val(r, f) for f in ("algorithm", "variant", "dataset", "N", "K", "seed", "mode", "rep", "run_id"))
        for r in rows
    )
    for key, count in key_counter.items():
        if count > 1:
            duplicate_keys.append({"key": key, "count": count})
    checks["duplicate_row_keys"] = {"ok": not duplicate_keys, "duplicates": duplicate_keys}

    checks["metadata"] = metadata_summary()
    checks["metadata_alignment"] = metadata_alignment(rows, checks["metadata"])
    aux_errors = []
    for line, row in enumerate(rows, start=2):
        if val(row, "mode") == "time":
            expected = max(0, int(val(row, "heap_peak_bytes")) - int(val(row, "heap_before_bytes")))
            actual = int(val(row, "aux_bytes"))
            if actual != expected:
                aux_errors.append({"line": line, "expected": expected, "actual": actual})
        elif val(row, "aux_bytes"):
            aux_errors.append({"line": line, "expected": "", "actual": val(row, "aux_bytes")})
    checks["aux_bytes_formula"] = {"ok": not aux_errors, "errors": aux_errors}
    checks["K_blank_for_d1_d2"] = all(not val(row, "K") for row in rows)
    checks["environment_values"] = sorted({
        (val(row, "jdk"), val(row, "os"), val(row, "arch"), val(row, "cores"),
         val(row, "max_heap_bytes")) for row in rows
    })
    checks["row_count_total"] = len(rows)
    checks["row_count_by_mode"] = dict(Counter(val(r, "mode") for r in rows))
    checks["row_count_by_dataset_mode"] = {
        f"{dataset}/{mode}": count
        for (dataset, mode), count in Counter((val(r, "dataset"), val(r, "mode")) for r in rows).items()
    }

    AUDIT_DIR.mkdir(parents=True, exist_ok=True)
    (AUDIT_DIR / "audit_spec_csv_result.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0 if all([
        checks["header_exact_code_order"],
        checks["minimum_required_fields_in_order"],
        checks["count_fields_present"],
        checks["csv_header_matches_expected_code_header"],
        checks["mode_separation"]["ok"],
        checks["sorted_ok_all_true"]["ok"],
        checks["output_check_time_count_equality"]["all_groups_equal"],
        checks["independent_d1_reproducibility"]["all_match"],
        checks["duplicate_row_keys"]["ok"],
    ]) else 1


if __name__ == "__main__":
    sys.exit(main())
